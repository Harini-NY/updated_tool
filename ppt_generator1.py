import copy
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

# ------------------------------------------------------------------ settings
FONT_NAME = "Aptos"          # same as your template's theme font
TABLE_FONT = 10              # STANDARD size (pt) for all body text in both tables and the takeaway
MIN_FONT = 7                 # if text can't fit at TABLE_FONT, ALL text shrinks together, never below this
SUMMARY_FONT = 12            # the two small "completed? / reviewed?" value cells (matches their label cells)
MIN_ROW_H = 12               # pt, floor for a row (real height comes from the text)
CELL_MARGIN_X = 6            # pt
CELL_MARGIN_Y = 3            # pt
START_ROW = 1                # row 0 is the header
REMOVE_UNUSED_ROWS = True    # fewer items than template rows -> rows are removed and the rest are made taller
MITIGATION_BULLETS = True    # bullet (like your template) when there is real mitigation text
BULLET_INDENT = 171450       # EMU, same as the template's bullet indent
EMU_PER_PT = 12700

IMPACT_COLOURS = {
    "+": RGBColor(0, 176, 80),       # green
    "=": RGBColor(255, 192, 0),      # yellow
    "-": RGBColor(237, 125, 49),     # orange
    "- -": RGBColor(255, 0, 0),      # red
    "NA": RGBColor(191, 191, 191),   # grey
}
BLACK = RGBColor(0, 0, 0)


# ------------------------------------------------------------------ text helpers
def fmt_value(value):
    """None / N/A -> NA (as in your template), 150.0 -> 150, 1234.5 -> 1,234.50"""
    if value is None:
        return "NA"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{int(value):,}" if float(value).is_integer() else f"{value:,.2f}"
    text = str(value).strip()
    return "NA" if text.upper() in ("N/A", "NA") else text


def count_lines(text, width_pt, font_pt, margin_x=CELL_MARGIN_X):
    """Estimate wrapped line count (greedy word wrap, ~0.55 x font size per character)."""
    per_line = max(1, int((width_pt - 2 * margin_x) / (font_pt * 0.55)))
    total = 0
    for para in str(text).split("\n"):
        lines, cur = 1, 0
        for word in para.split():
            for i in range(0, len(word), per_line):          # break very long words
                chunk = word[i:i + per_line]
                add = len(chunk) + (1 if cur else 0)
                if cur + add > per_line:
                    lines += 1
                    cur = len(chunk)
                else:
                    cur += add
        total += lines
    return total


def needed_height(text, width_pt, font_pt):
    return count_lines(text, width_pt, font_pt) * font_pt * 1.2 + 2 * CELL_MARGIN_Y


def set_impact_colour(cell, impact):
    cell.fill.solid()
    cell.fill.fore_color.rgb = IMPACT_COLOURS.get(str(impact).strip().upper(), RGBColor(255, 255, 255))


def set_bullet(paragraph, on):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buClr", "a:buFont", "a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if on:
        pPr.set("marL", str(BULLET_INDENT))
        pPr.set("indent", str(-BULLET_INDENT))
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "\u2022")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))


def format_cell(cell, text, font_size, bold=False, align=PP_ALIGN.CENTER, bullet=False):
    cell.text = str(text)
    tf = cell.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    cell.margin_left = cell.margin_right = Pt(CELL_MARGIN_X)
    cell.margin_top = cell.margin_bottom = Pt(CELL_MARGIN_Y)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        set_bullet(paragraph, bullet)
        # Empty cells have no runs, so PowerPoint would fall back to a big default size
        paragraph._p.get_or_add_endParaRPr().set("sz", str(int(font_size * 100)))
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = BLACK


# ------------------------------------------------------------------ table layout
def name_parameter(it):
    unit = str(it.get("Unit") or "").strip()
    return f"{str(it['Parameter']).strip()} ({unit})" if unit else str(it["Parameter"]).strip()


def name_system(it):
    unit = str(it.get("Unit") or "").strip()
    return f"{str(it['System']).strip()}, {unit}" if unit else str(it["System"]).strip()


def table_columns(name_fn, limit_key):
    """(value getter, alignment, is_impact, is_mitigation) for the 7 columns."""
    return [
        (name_fn, PP_ALIGN.LEFT, False, False),
        (lambda it: it["Baseline"], PP_ALIGN.CENTER, False, False),
        (lambda it: it["Projected Load"], PP_ALIGN.CENTER, False, False),
        (lambda it: it[limit_key], PP_ALIGN.CENTER, False, False),
        (lambda it: it["Impact"], PP_ALIGN.CENTER, True, False),
        (lambda it: it["Mitigation"], PP_ALIGN.LEFT, False, True),
        (lambda it: it["New Impact"], PP_ALIGN.CENTER, True, False),
    ]


def has_bullet(text, is_mitigation):
    return MITIGATION_BULLETS and is_mitigation and text not in ("", "NA")


def fit_header(shape):
    """The template's header rows are a little shorter than their own wrapped text, so
    PowerPoint grows them and pushes the table into the one below. Size the header to its
    text and return the height left for the data rows (total table height stays the same)."""
    table = shape.table
    total = sum(r.height for r in table.rows) / EMU_PER_PT
    need = 0
    for c, cell in enumerate(table.rows[0].cells):
        sizes = [r.font.size.pt for p in cell.text_frame.paragraphs for r in p.runs if r.font.size]
        size = max(sizes) if sizes else 10
        lines = count_lines(cell.text_frame.text, table.columns[c].width / EMU_PER_PT, size,
                            margin_x=cell.margin_left / EMU_PER_PT)
        need = max(need, lines * size * 1.2 + (cell.margin_top + cell.margin_bottom) / EMU_PER_PT)
    header = max(table.rows[0].height / EMU_PER_PT, need)
    table.rows[0].height = Pt(header)
    return total - header


def row_heights(job, font):
    """Height (pt) each item's row needs at this font size."""
    heights = []
    for item in job["items"]:
        h = MIN_ROW_H
        for c, (getter, _align, _imp, is_mit) in enumerate(job["columns"]):
            text = fmt_value(getter(item))
            width = job["widths"][c] - (BULLET_INDENT / EMU_PER_PT if has_bullet(text, is_mit) else 0)
            h = max(h, needed_height(text, width, font))
        heights.append(h)
    return heights


def standard_font(jobs):
    """One font size for the whole slide: TABLE_FONT, or the largest smaller size at which
    both tables still fit inside the space the template gives them."""
    size = TABLE_FONT
    while size > MIN_FONT and not all(sum(row_heights(j, size)) <= j["budget"] for j in jobs):
        size -= 0.5
    return size


def fill_table(job, font):
    shape, items, columns = job["shape"], job["items"], job["columns"]
    table = shape.table
    tbl = table._tbl

    n_data = max(len(items), 1)
    if not REMOVE_UNUSED_ROWS:
        n_data = max(n_data, len(table.rows) - START_ROW)
    needed = START_ROW + n_data
    while len(table.rows) < needed:                    # more items than template rows: clone last row
        clone = copy.deepcopy(tbl.tr_lst[-1])
        for ext in clone.findall(qn("a:extLst")):      # drop the row id so ids stay unique
            clone.remove(ext)
        tbl.append(clone)
    while len(table.rows) > needed:
        tbl.remove(tbl.tr_lst[-1])

    heights = row_heights(job, font) + [MIN_ROW_H] * (n_data - len(items))
    extra = max(0.0, job["budget"] - sum(heights)) / n_data     # spread spare space so the table keeps its frame
    if sum(heights) > job["budget"] + 0.5:
        print(f"WARNING: {shape.name} needs more height than the template gives it; "
              f"too much text for one slide.")

    for r in range(START_ROW, needed):
        idx = r - START_ROW
        item = items[idx] if idx < len(items) else None
        for c, (getter, align, is_impact, is_mit) in enumerate(columns):
            text = fmt_value(getter(item)) if item else ""
            cell = table.cell(r, c)
            format_cell(cell, text, font, bold=is_impact, align=align,
                        bullet=has_bullet(text, is_mit))
            if is_impact:
                set_impact_colour(cell, text)
        table.rows[r].height = Pt(heights[idx] + extra)

    shape.height = int(sum(row.height for row in table.rows))


def set_takeaway(shape, text, font_size):
    """Heading + body at the standard size (shrinks only if it cannot fit the box)."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    width = shape.width / EMU_PER_PT - (tf.margin_left + tf.margin_right) / EMU_PER_PT
    height = shape.height / EMU_PER_PT - (tf.margin_top + tf.margin_bottom) / EMU_PER_PT
    body = str(text).strip()
    size = font_size
    while size > MIN_FONT and (count_lines(body, width, size, margin_x=0) + 1) * size * 1.2 > height:
        size -= 0.5

    tf.clear()
    lines = [("Key Takeaway", True)] + [(ln, False) for ln in body.split("\n")] if body else [("Key Takeaway", True)]
    for i, (line, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT_NAME, Pt(size), bold, BLACK


# ------------------------------------------------------------------ main
def generate_ppt(parameter_results, systems_results, tool_completed,
                 design_reviewed, key_takeaway, site, project_title):
    prs = Presentation("Templates/template.pptx")
    slide = prs.slides[0]

    # Find the tables by shape: two big (7 columns) ones stacked, two small (2 columns) ones side by side
    tables = [s for s in slide.shapes if s.has_table]
    big = sorted([s for s in tables if len(s.table.columns) >= 7], key=lambda s: s.top)
    small = sorted([s for s in tables if len(s.table.columns) == 2], key=lambda s: s.left)
    if len(big) < 2 or len(small) < 2:
        raise ValueError("Template slide 1 needs 2 big tables (7 columns) and 2 small tables (2 columns).")
    parameter_shape, systems_shape = big[0], big[1]
    assessment_shape, design_shape = small[0], small[1]

    jobs = []
    for shape, items, columns in (
        (parameter_shape, parameter_results, table_columns(name_parameter, "Tariff")),
        (systems_shape, systems_results, table_columns(name_system, "Redundancy")),
    ):
        jobs.append({
            "shape": shape, "items": items, "columns": columns,
            "widths": [c.width / EMU_PER_PT for c in shape.table.columns],
            # space left for the data rows: the table stays inside its frame, so nothing moves or overlaps
            "budget": fit_header(shape),
        })

    font = standard_font(jobs)       # one size for everything on the slide
    for job in jobs:
        fill_table(job, font)

    # Summary cells: same 12pt as their label cells
    format_cell(assessment_shape.table.cell(0, 1), fmt_value(tool_completed), SUMMARY_FONT, align=PP_ALIGN.LEFT)
    format_cell(design_shape.table.cell(0, 1), fmt_value(design_reviewed), SUMMARY_FONT, align=PP_ALIGN.LEFT)

    # Key takeaway box
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().startswith("Key Takeaway"):
            set_takeaway(shape, key_takeaway, font)
            break

    os.makedirs("saved_assessments", exist_ok=True)
    safe_title = re.sub(r'[\\/:*?"<>|]', "_", str(project_title))
    output_file = f"saved_assessments/{safe_title}.pptx"
    prs.save(output_file)
    return output_file

